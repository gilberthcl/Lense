"""Smoke tests for DOCX report generation."""
import io

from docx import Document

from app.services import report_docx


def _sample_findings():
    return [
        {
            "id": 1, "finding_ref": "F-001", "title": "Suspicious PowerShell",
            "category": "suspicious", "severity": "high", "confidence": "medium",
            "status": "validated", "summary": "Encoded command observed.",
            "evidence": {"verbatim_values": ["10.0.0.9"]},
            "mitre": [{"technique_id": "T1059.001", "name": "PowerShell"}],
            "affected_assets": ["WKSTN-01"], "affected_users": ["jdoe"],
            "recommendations": "Isolate host.", "dataset_id": 1,
        },
        {
            "id": 2, "finding_ref": "F-002", "title": "Same host, second dataset",
            "category": "malicious", "severity": "critical", "confidence": "high",
            "status": "draft", "summary": "Lateral movement.",
            "evidence": {"rows": ["10.0.0.9 -> WKSTN-01"]},
            "mitre": [{"technique_id": "T1021", "name": "Remote Services"}],
            "affected_assets": ["WKSTN-01"], "affected_users": [],
            "recommendations": "Investigate.", "dataset_id": 2,
        },
    ]


def _build(lang):
    return report_docx.build_report(
        tenant_name="Acme Corp",
        hunt={"name": "Q2 Hunt", "objective": "Find lateral movement"},
        findings=_sample_findings(),
        datasets=[
            {"dataset_id": 1, "filename": "logons.csv", "assessment": "Clean-ish."},
            {"dataset_id": 2, "filename": "proc.csv", "assessment": "Notable."},
        ],
        generated_at="2026-06-22 12:00 UTC",
        lang=lang,
    )


def test_report_is_valid_docx_english():
    data = _build("en")
    assert data[:2] == b"PK"  # docx is a zip
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Threat Hunt Findings Report" in text
    assert "Acme Corp" in text
    # findings + IOC + MITRE + correlation tables all rendered
    assert len(doc.tables) >= 4


def test_report_spanish_labels():
    data = _build("es")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Informe de Hallazgos de Threat Hunting" in text
    assert "Resumen Ejecutivo" in text


def test_report_handles_no_findings():
    data = report_docx.build_report(
        tenant_name="Acme", hunt={"name": "Empty", "objective": None},
        findings=[], datasets=[{"dataset_id": 1, "filename": "x.csv", "assessment": None}],
        generated_at="2026-06-22 12:00 UTC", lang="en",
    )
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "No findings were identified" in text
