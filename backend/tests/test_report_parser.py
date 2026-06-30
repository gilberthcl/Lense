"""Report parser — heading-driven sectioning of a hunt report .docx."""
import io

from docx import Document

from app.services import report_parser


def _build_report() -> bytes:
    d = Document()
    d.add_heading("Executive Summary", level=1)
    d.add_paragraph("Between Jan 3 and Jan 31 the team searched for actor X.")
    d.add_heading("Background", level=1)
    d.add_paragraph("Actor X is a destructive group.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Technique ID & Name"; t.cell(0, 1).text = "Tactic"
    t.cell(1, 0).text = "T1003.001 – LSASS Memory"; t.cell(1, 1).text = "Credential Access"
    d.add_heading("Findings", level=1)
    d.add_heading("LSASS Credential Dump on Domain Controller", level=3)
    d.add_paragraph("rundll32.exe comsvcs.dll MiniDump 680 lsass.dmp full")
    d.add_paragraph("This dumps LSASS for offline credential extraction.")
    d.add_heading("ADRecon Enumeration from Workstation", level=3)
    d.add_paragraph("powershell.exe -File dra.ps1 ran AD recon.")
    d.add_heading("Methodology", level=1)
    d.add_paragraph("Action Plan: hunt for credential theft then lateral movement.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_parse_extracts_sections_findings_methodology_mitre():
    r = report_parser.parse_report_docx(_build_report())

    assert "Executive Summary" in r["section_names"]
    assert "Findings" in r["section_names"]

    # One entry per finding heading, each with its body, no duplicates.
    titles = [f["title"] for f in r["findings"]]
    assert titles == [
        "LSASS Credential Dump on Domain Controller",
        "ADRecon Enumeration from Workstation",
    ]
    assert "comsvcs.dll" in r["findings"][0]["body"]
    assert "AD recon" in r["findings"][1]["body"]

    assert "Methodology" in r["methodology_text"]
    assert "Action Plan" in r["methodology_text"]

    assert r["mitre_table"] and "T1003.001" in r["mitre_table"]


def test_findings_outside_a_findings_section_are_not_captured():
    d = Document()
    d.add_heading("Background", level=1)
    d.add_heading("Some Subsection", level=3)   # a deep heading, but NOT under Findings
    d.add_paragraph("context text")
    buf = io.BytesIO()
    d.save(buf)
    r = report_parser.parse_report_docx(buf.getvalue())
    assert r["findings"] == []  # nothing mis-detected as a finding


def test_empty_doc_is_safe():
    d = Document()
    buf = io.BytesIO()
    d.save(buf)
    r = report_parser.parse_report_docx(buf.getvalue())
    assert r["findings"] == [] and r["mitre_table"] is None
