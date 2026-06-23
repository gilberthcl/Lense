"""Tests for methodology doc extraction and prompt-template wiring."""
import io

import pytest
from docx import Document

from app.services import doc_loader, prompts


def test_extract_text_txt_and_md():
    assert doc_loader.extract_text("m.txt", b"hello world") == "hello world"
    assert "# Plan" in doc_loader.extract_text("m.md", b"# Plan\nstep 1")


def test_extract_text_rejects_unknown():
    with pytest.raises(ValueError):
        doc_loader.extract_text("m.pdf", b"%PDF")


def test_extract_docx_includes_tables():
    doc = Document()
    doc.add_paragraph("Plan of Action")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Topic 1"
    table.rows[0].cells[1].text = "SMB"
    table.rows[1].cells[0].text = "Topic 2"
    table.rows[1].cells[1].text = "RDP"
    buf = io.BytesIO()
    doc.save(buf)
    text = doc_loader.extract_text("m.docx", buf.getvalue())
    assert "Plan of Action" in text
    assert "Topic 1 | SMB" in text  # table content preserved
    assert "[TABLE]" in text


def test_analyst_prompts_format_with_all_fields():
    # Ensures every placeholder is supplied — guards against KeyError at runtime.
    sys = prompts.ANALYST_SYSTEM.format(
        analysis_instructions="proto", guardrails="g", categories="cats"
    )
    assert "proto" in sys and "cats" in sys
    user = prompts.ANALYST_PROMPT.format(
        hunt_name="H", edr="Falcon", siem="QRadar", language="Spanish",
        dataset_focus="focus", methodology_brief="brief", methodology="meth",
        tenant_context="ctx", dataset_name="d.csv", evidence_json="{}",
    )
    assert "Spanish" in user and "Falcon" in user and "QRadar" in user


def test_methodology_prompt_formats():
    out = prompts.METHODOLOGY_PROMPT.format(
        edr="Falcon", siem="QRadar", language="English", methodology="text"
    )
    assert "Falcon" in out and "text" in out
