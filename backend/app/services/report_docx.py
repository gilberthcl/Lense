"""
DOCX report generation (python-docx).

Produces a professional, report-ready threat-hunt findings document:
title block, executive summary (deterministic counts — no invented narrative),
per-dataset assessments, the full findings table, an IOC table, a MITRE ATT&CK
technique mapping, and cross-dataset correlations.

Bilingual: the operator writes client reports in Spanish and technical content
in English, so every label is available in both languages (`lang="es"|"en"`).

Evidence discipline applies here too: this module only renders values that the
findings pipeline already produced. It computes counts; it never asserts new
facts about hosts, users, or indicators.
"""
from __future__ import annotations

import io
from collections import Counter
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.services import correlation

# ── Bilingual label catalogue ───────────────────────────────────────────────
LABELS: dict[str, dict[str, str]] = {
    "report_title": {"en": "Threat Hunt Findings Report", "es": "Informe de Hallazgos de Threat Hunting"},
    "tenant": {"en": "Client", "es": "Cliente"},
    "hunt": {"en": "Hunt", "es": "Cacería"},
    "objective": {"en": "Objective", "es": "Objetivo"},
    "generated": {"en": "Generated", "es": "Generado"},
    "exec_summary": {"en": "Executive Summary", "es": "Resumen Ejecutivo"},
    "datasets_analyzed": {"en": "Datasets analyzed", "es": "Conjuntos de datos analizados"},
    "total_findings": {"en": "Total findings", "es": "Hallazgos totales"},
    "validated_findings": {"en": "Validated findings", "es": "Hallazgos validados"},
    "by_category": {"en": "Findings by category", "es": "Hallazgos por categoría"},
    "by_severity": {"en": "Findings by severity", "es": "Hallazgos por severidad"},
    "no_findings_note": {
        "en": "No findings were identified for this hunt. \"No Finding\" is a valid outcome.",
        "es": "No se identificaron hallazgos para esta cacería. \"Sin Hallazgo\" es un resultado válido.",
    },
    "dataset_assessments": {"en": "Per-Dataset Assessments", "es": "Evaluaciones por Conjunto de Datos"},
    "no_assessment": {"en": "No assessment recorded.", "es": "Sin evaluación registrada."},
    "findings_detail": {"en": "Findings", "es": "Hallazgos"},
    "ioc_table": {"en": "Indicators of Compromise", "es": "Indicadores de Compromiso"},
    "mitre_mapping": {"en": "MITRE ATT&CK Mapping", "es": "Mapeo MITRE ATT&CK"},
    "correlations": {"en": "Cross-Dataset Correlations", "es": "Correlaciones Entre Conjuntos de Datos"},
    "correlations_note": {
        "en": "Entities appearing in findings across multiple datasets — potential campaign signal.",
        "es": "Entidades presentes en hallazgos de múltiples conjuntos de datos — posible señal de campaña.",
    },
    "none": {"en": "None", "es": "Ninguno"},
    # table headers
    "ref": {"en": "Ref", "es": "Ref"},
    "title": {"en": "Title", "es": "Título"},
    "category": {"en": "Category", "es": "Categoría"},
    "severity": {"en": "Severity", "es": "Severidad"},
    "confidence": {"en": "Confidence", "es": "Confianza"},
    "status": {"en": "Status", "es": "Estado"},
    "summary": {"en": "Summary", "es": "Resumen"},
    "evidence": {"en": "Evidence", "es": "Evidencia"},
    "mitre": {"en": "MITRE", "es": "MITRE"},
    "assets": {"en": "Affected Assets", "es": "Activos Afectados"},
    "users": {"en": "Affected Users", "es": "Usuarios Afectados"},
    "recommendations": {"en": "Recommendations", "es": "Recomendaciones"},
    "indicator": {"en": "Indicator", "es": "Indicador"},
    "type": {"en": "Type", "es": "Tipo"},
    "occurrences": {"en": "Occurrences", "es": "Ocurrencias"},
    "technique": {"en": "Technique", "es": "Técnica"},
    "technique_id": {"en": "Technique ID", "es": "ID de Técnica"},
    "count": {"en": "Count", "es": "Conteo"},
    "entity": {"en": "Entity", "es": "Entidad"},
    "datasets": {"en": "Datasets", "es": "Conjuntos"},
    "findings_col": {"en": "Findings", "es": "Hallazgos"},
}

CATEGORY_ORDER = ["malicious", "suspicious", "risky", "policy_violation", "unconfirmed", "no_finding"]
SEVERITY_ORDER = ["critical", "high", "medium", "low", "info"]

_HEADER_FILL = "1F2937"  # slate-800-ish header shading


def _t(key: str, lang: str) -> str:
    entry = LABELS.get(key, {})
    return entry.get(lang) or entry.get("en") or key


def _as_lines(value: Any) -> str:
    return ", ".join(correlation._as_str_list(value)) or "—"


def _normalize_mitre(value: Any) -> list[tuple[str, str]]:
    """Return [(technique_id, name)] from a finding's mitre field (defensive)."""
    out: list[tuple[str, str]] = []
    if isinstance(value, dict):
        value = [value]
    if isinstance(value, list):
        for item in value:
            if isinstance(item, dict):
                tid = str(item.get("technique_id") or item.get("id") or "").strip()
                name = str(item.get("name") or item.get("technique") or "").strip()
                if tid or name:
                    out.append((tid or "—", name or "—"))
            elif item:
                out.append((str(item), "—"))
    elif isinstance(value, str) and value.strip():
        out.append((value.strip(), "—"))
    return out


# ── Low-level docx helpers ──────────────────────────────────────────────────
def _shade_cell(cell, fill: str) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _header_row(table, headers: list[str]) -> None:
    row = table.rows[0]
    for i, text in enumerate(headers):
        cell = row.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(cell, _HEADER_FILL)


def _add_table(doc, headers: list[str]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _header_row(table, headers)
    return table


def _set_cell(cell, text: str, size: int = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text if text else "—")
    run.font.size = Pt(size)


def _heading(doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


# ── Section builders ────────────────────────────────────────────────────────
def _build_title(doc, *, tenant_name: str, hunt: dict, generated_at: str, lang: str) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(_t("report_title", lang))
    run.bold = True
    run.font.size = Pt(22)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{_t('tenant', lang)}: {tenant_name}").bold = True

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"{_t('hunt', lang)}: {hunt.get('name', '—')}    |    "
        f"{_t('generated', lang)}: {generated_at}"
    )
    if hunt.get("objective"):
        obj = doc.add_paragraph()
        obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
        obj.add_run(f"{_t('objective', lang)}: {hunt['objective']}").italic = True


def _build_exec_summary(doc, *, findings: list[dict], dataset_count: int, lang: str) -> None:
    _heading(doc, _t("exec_summary", lang), level=1)

    real = [f for f in findings if f.get("category") != "no_finding"]
    validated = sum(1 for f in real if f.get("status") == "validated")
    doc.add_paragraph(f"{_t('datasets_analyzed', lang)}: {dataset_count}")
    doc.add_paragraph(f"{_t('total_findings', lang)}: {len(real)}")
    doc.add_paragraph(f"{_t('validated_findings', lang)}: {validated}")

    if not real:
        p = doc.add_paragraph()
        p.add_run(_t("no_findings_note", lang)).italic = True
        return

    cat_counts = Counter(f.get("category", "unconfirmed") for f in real)
    sev_counts = Counter((f.get("severity") or "info").lower() for f in real)

    doc.add_paragraph(_t("by_category", lang), style="List Bullet").runs[0].bold = True
    for cat in CATEGORY_ORDER:
        if cat_counts.get(cat):
            doc.add_paragraph(f"{cat.replace('_', ' ')}: {cat_counts[cat]}", style="List Bullet 2")

    doc.add_paragraph(_t("by_severity", lang), style="List Bullet").runs[0].bold = True
    for sev in SEVERITY_ORDER:
        if sev_counts.get(sev):
            doc.add_paragraph(f"{sev}: {sev_counts[sev]}", style="List Bullet 2")


def _build_dataset_assessments(doc, *, datasets: list[dict], lang: str) -> None:
    if not datasets:
        return
    _heading(doc, _t("dataset_assessments", lang), level=1)
    for ds in datasets:
        doc.add_paragraph(ds.get("filename", "—"), style="List Bullet").runs[0].bold = True
        assessment = (ds.get("assessment") or "").strip() or _t("no_assessment", lang)
        doc.add_paragraph(assessment, style="List Bullet 2")


def _build_findings(doc, *, findings: list[dict], lang: str) -> None:
    real = [f for f in findings if f.get("category") != "no_finding"]
    _heading(doc, _t("findings_detail", lang), level=1)
    if not real:
        doc.add_paragraph(_t("none", lang))
        return

    headers = [_t("ref", lang), _t("title", lang), _t("category", lang),
               _t("severity", lang), _t("confidence", lang), _t("status", lang)]
    table = _add_table(doc, headers)
    for f in real:
        cells = table.add_row().cells
        _set_cell(cells[0], str(f.get("finding_ref") or "—"))
        _set_cell(cells[1], str(f.get("title") or "—"))
        _set_cell(cells[2], str(f.get("category") or "—").replace("_", " "))
        _set_cell(cells[3], str(f.get("severity") or "—"))
        _set_cell(cells[4], str(f.get("confidence") or "—"))
        _set_cell(cells[5], str(f.get("status") or "—"))

    # Detailed per-finding blocks beneath the summary table.
    for f in real:
        doc.add_paragraph()
        h = doc.add_paragraph()
        run = h.add_run(f"{f.get('finding_ref', '—')} — {f.get('title', '—')}")
        run.bold = True
        run.font.size = Pt(12)
        for key in ("category", "severity", "confidence", "status"):
            doc.add_paragraph(f"{_t(key, lang)}: {f.get(key) or '—'}")
        if f.get("summary"):
            doc.add_paragraph(f"{_t('summary', lang)}: {f['summary']}")
        doc.add_paragraph(f"{_t('assets', lang)}: {_as_lines(f.get('affected_assets'))}")
        doc.add_paragraph(f"{_t('users', lang)}: {_as_lines(f.get('affected_users'))}")
        mitre = "; ".join(f"{tid} {name}".strip() for tid, name in _normalize_mitre(f.get("mitre")))
        doc.add_paragraph(f"{_t('mitre', lang)}: {mitre or '—'}")
        if f.get("recommendations"):
            doc.add_paragraph(f"{_t('recommendations', lang)}: {_as_lines(f.get('recommendations'))}")


def _build_iocs(doc, *, iocs: list[dict], lang: str) -> None:
    _heading(doc, _t("ioc_table", lang), level=1)
    if not iocs:
        doc.add_paragraph(_t("none", lang))
        return
    table = _add_table(doc, [_t("indicator", lang), _t("type", lang), _t("occurrences", lang)])
    for ioc in iocs:
        cells = table.add_row().cells
        _set_cell(cells[0], str(ioc.get("value")))
        _set_cell(cells[1], str(ioc.get("entity_type")))
        _set_cell(cells[2], str(ioc.get("finding_count")))


def _build_mitre(doc, *, findings: list[dict], lang: str) -> None:
    _heading(doc, _t("mitre_mapping", lang), level=1)
    counts: Counter = Counter()
    names: dict[str, str] = {}
    for f in findings:
        for tid, name in _normalize_mitre(f.get("mitre")):
            counts[tid] += 1
            if name and name != "—":
                names[tid] = name
    if not counts:
        doc.add_paragraph(_t("none", lang))
        return
    table = _add_table(doc, [_t("technique_id", lang), _t("technique", lang), _t("count", lang)])
    for tid, count in counts.most_common():
        cells = table.add_row().cells
        _set_cell(cells[0], tid)
        _set_cell(cells[1], names.get(tid, "—"))
        _set_cell(cells[2], str(count))


def _build_correlations(doc, *, correlations: list[dict], lang: str) -> None:
    _heading(doc, _t("correlations", lang), level=1)
    doc.add_paragraph(_t("correlations_note", lang)).runs[0].italic = True
    if not correlations:
        doc.add_paragraph(_t("none", lang))
        return
    table = _add_table(doc, [_t("entity", lang), _t("type", lang),
                             _t("datasets", lang), _t("findings_col", lang), _t("category", lang)])
    for c in correlations:
        cells = table.add_row().cells
        _set_cell(cells[0], str(c.get("value")))
        _set_cell(cells[1], str(c.get("entity_type")))
        _set_cell(cells[2], str(c.get("dataset_count")))
        _set_cell(cells[3], str(c.get("finding_count")))
        _set_cell(cells[4], str(c.get("max_category", "—")).replace("_", " "))


def build_report(
    *,
    tenant_name: str,
    hunt: dict[str, Any],
    findings: list[dict[str, Any]],
    datasets: list[dict[str, Any]],
    generated_at: str,
    lang: str = "en",
) -> bytes:
    """
    Assemble the full DOCX report and return it as bytes.

    `datasets` items: {dataset_id, filename, assessment?}
    `findings`  items: finding dicts (see Finding model).
    """
    lang = "es" if str(lang).lower().startswith("es") else "en"
    dataset_names = {d["dataset_id"]: d.get("filename", "") for d in datasets if d.get("dataset_id") is not None}
    corr = correlation.compute_correlations(findings, dataset_names)

    doc = Document()
    _build_title(doc, tenant_name=tenant_name, hunt=hunt, generated_at=generated_at, lang=lang)
    _build_exec_summary(doc, findings=findings, dataset_count=len(datasets), lang=lang)
    _build_dataset_assessments(doc, datasets=datasets, lang=lang)
    _build_findings(doc, findings=findings, lang=lang)
    _build_iocs(doc, iocs=corr["iocs"], lang=lang)
    _build_mitre(doc, findings=findings, lang=lang)
    _build_correlations(doc, correlations=corr["correlations"], lang=lang)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
